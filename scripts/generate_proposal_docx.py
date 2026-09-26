import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_proposal_docx(output_path):
    doc = Document()
    
    # Sayfa Kenar Boşlukları (Normal)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Renk Paleti (Kurumsal Lacivert / İnşaat Mühendisliği)
    PRIMARY_COLOR = RGBColor(15, 23, 42)      # Slate 900
    ACCENT_COLOR = RGBColor(2, 132, 199)      # Sky 600
    TEXT_MUTED = RGBColor(100, 116, 139)      # Slate 500
    TEXT_DARK = RGBColor(30, 41, 59)          # Slate 800
    BORDER_COLOR = "CBD5E1"
    
    # 1. BAŞLIK VE ÜST BİLGİ
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.5)
    header_table.columns[1].width = Inches(2.4)
    
    # Sol Hücre (Başlık & Slogan)
    c0 = header_table.cell(0, 0)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    run_main = p0.add_run("RENOVATION OS")
    run_main.font.name = "Calibri"
    run_main.font.size = Pt(24)
    run_main.font.bold = True
    run_main.font.color.rgb = PRIMARY_COLOR
    
    p0_sub = c0.add_paragraph()
    p0_sub.paragraph_format.space_after = Pt(0)
    run_sub = p0_sub.add_run("Proje, Muhasebe ve Müşteri Yönetim Sistemi")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = ACCENT_COLOR
    
    p0_desc = c0.add_paragraph()
    p0_desc.paragraph_format.space_after = Pt(0)
    run_desc = p0_desc.add_run("Yazılım Mimarisi, Fonksiyonel Şartname ve Ticari Teklif")
    run_desc.font.name = "Calibri"
    run_desc.font.size = Pt(9.5)
    run_desc.font.color.rgb = TEXT_MUTED
    
    # Sağ Hücre (Teklif Kartı / Fiyat Kutusu)
    c1 = header_table.cell(0, 1)
    set_cell_background(c1, "F8FAFC")
    set_cell_margins(c1, 140, 140, 180, 180)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_after = Pt(2)
    
    r_date = p1.add_run("Tarih: 25.09.2026\nKod: REN-OS-2026-V1\n")
    r_date.font.size = Pt(8.5)
    r_date.font.color.rgb = TEXT_MUTED
    
    r_price_label = p1.add_run("Teklif Tutarı:\n")
    r_price_label.font.size = Pt(9)
    r_price_label.font.bold = True
    r_price_label.font.color.rgb = PRIMARY_COLOR
    
    r_price = p1.add_run("1.000 USD\n")
    r_price.font.size = Pt(16)
    r_price.font.bold = True
    r_price.font.color.rgb = ACCENT_COLOR
    
    r_price_sub = p1.add_run("Anahtar Teslim • Sıfır Aylık Ücret")
    r_price_sub.font.size = Pt(8)
    r_price_sub.font.italic = True
    r_price_sub.font.color.rgb = TEXT_MUTED

    # Çizgi Ayırıcı
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(12)
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("―" * 58)
    run_div.font.color.rgb = RGBColor(226, 232, 240)

    # 2. YÖNETİCİ ÖZETİ
    h1 = doc.add_heading("1. YÖNETİCİ ÖZETİ VE ÇÖZÜM VİZYONU", level=1)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR

    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.space_after = Pt(8)
    p_summary.paragraph_format.line_spacing = 1.15
    r_s = p_summary.add_run(
        "İnşaat ve renovasyon sektörünün sahadaki dinamikleri (değişken keşifler, parçalı malzeme alımları, usta yevmiyeleri, "
        "aşamalı hakediş tahsilatları ve küçük çaplı malzeme alım-satım ticareti) standart piyasa muhasebe programlarının kalıplarına uymaz.\n\n"
        "Size özel olarak mimarisi kurgulanan RENOVATION OS, şirketinizin müşteri adayından keşif randevusuna, detaylı malzeme/işçilik teklifinden "
        "resmi sözleşmeye, şantiye harcamalarından usta puantajına, kasa/banka takibinden net şirket kârlılığına ve ürün alım-satım stok hareketlerine kadar "
        "tüm operasyonunu tek bir merkezden, yalın, yüksek hızlı ve hatasız yönetmenizi sağlar."
    )
    r_s.font.size = Pt(10)
    r_s.font.color.rgb = TEXT_DARK

    # Önemli Vurgu Kutusu (Callout Box)
    callout_table = doc.add_table(rows=1, cols=1)
    callout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_table.columns[0].width = Inches(6.9)
    c_box = callout_table.cell(0, 0)
    set_cell_background(c_box, "EFF6FF")
    set_cell_margins(c_box, 140, 140, 180, 180)
    p_box = c_box.paragraphs[0]
    p_box.paragraph_format.space_after = Pt(0)
    p_box.paragraph_format.line_spacing = 1.15
    r_box_title = p_box.add_run("EN HESAPLI VE ÖZGÜR KULLANIM GARANTİSİ (SIFIR BAĞIMLILIK):\n")
    r_box_title.font.bold = True
    r_box_title.font.size = Pt(9.5)
    r_box_title.font.color.rgb = RGBColor(30, 64, 175)
    r_box_body = p_box.add_run(
        "Sistem tamamlanıp tarafınıza teslim edildikten sonra normal kullanım için harici bir yazılımcıya, sistem yöneticisine veya aylık sunucu "
        "bakım ücretlerine olan bağımlılığınız SIFIR olacaktır. Sistem tek tıkla cPanel veya standart herhangi bir web hostinge atılıp çalışır."
    )
    r_box_body.font.size = Pt(9.5)
    r_box_body.font.color.rgb = RGBColor(30, 58, 138)

    # 3. FONKSİYONEL SİSTEM MODÜLLERİ
    h2 = doc.add_heading("2. A'DAN Z'YE SİSTEM MODÜLLERİ VE ÖZELLİKLER", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR

    modules = [
        ("MODÜL 1: Müşteri Takip Sistemi (CRM) & Keşif Ajandası", 
         "• Müşteri Kartları: İletişim bilgileri, şantiye adresi ve özel notlar tek ekranda.\n"
         "• Aşama Yönetimi: Yeni Talep ➔ Keşif ➔ Teklif Verildi ➔ Anlaşıldı ➔ Şantiye Devam Ediyor ➔ Tamamlandı.\n"
         "• Randevu Takvimi: Keşif ve müşteri görüşme saatleri ajanda üzerinde izlenir; yaklaşan işler anasayfada hatırlatılır.\n"
         "• Müşteri Geçmişi: Müşterinin eski teklifleri, geçmiş ödemeleri ve şantiyeleri tek tıkla listelenir."),
        
        ("MODÜL 2: Teklif Hazırlama, Revizyon & Resmi Sözleşme",
         "• Dinamik Kalem Girişi: Teklif oluştururken satır satır Malzeme, İşçilik ve Nakliye kalemleri ekleme.\n"
         "• Canlı Hesaplama: Birim fiyat, miktar, iskonto ve KDV girildikçe toplam tutar otomatik hesaplanır.\n"
         "• 1-Tık Resmi PDF / Yazdırma: Firmanızın logosu ve antetiyle, müşteriye doğrudan WhatsApp veya E-Posta ile iletilebilir şık teklif/sözleşme belgesi.\n"
         "• Otomatik Projeye Dönüştürme: Müşteri teklifi onayladığında tek tuşla yeni bir 'Aktif Şantiye' başlatılır."),

        ("MODÜL 3: Şantiye & Proje Yönetimi (Maliyet & İlerleme)",
         "• Şantiye Kartı: Başlangıç ve teslim tarihi, sözleşme tutarı ve yetkili personel takibi.\n"
         "• Hakediş ve Ödeme Planı: Aşamalı ödeme takvimi (%30 Peşinat, %40 Kaba Bitimi vb.); ödenen ve kalan bakiye anlık izlenir.\n"
         "• Şantiye Harcamaları: Alınan çimento, boya, seramik fiş/faturaları doğrudan ilgili şantiyeye gider kaydedilir.\n"
         "• Canlı Şantiye Kârlılığı: Toplam tahsilat - Yapılan harcamalar = Kalan net kâr oranı anlık grafik olarak izlenir."),

        ("MODÜL 4: Usta & İşçilik Maliyetleri Takibi",
         "• 3 Esnek Çalışma Modeli: Yevmiye Usulü (Gün × Yevmiye), Götürü/Taşeron Usulü (Komple İş Bedeli) ve Metrekare/Birim Usulü.\n"
         "• Usta Cari Hesabı: Hangi ustaya ne kadar tahakkuk etti, ne kadarı elden/bankadan ödendi, kalan borç nedir eksiksiz takip edilir.\n"
         "• Puantaj ve Avans Takibi: Ustalara verilen avanslar ve hakediş ödemeleri tek tabloda listelenir."),

        ("MODÜL 5: Kasa Yönetimi (Nakit & Banka) ve Genel Giderler",
         "• Çift Kasa Sistemi: Nakit Kasa (Ofis/Elden) ve Banka Hesapları (Havale/EFT/Kredi Kartı) anlık bakiyeleri.\n"
         "• Şirket Sabit Giderleri: Projelerden bağımsız olan ofis kirası, şirket araçlarının akaryakıt/bakımları, yemek ve faturalar.\n"
         "• Belge & Makbuz Düzeni: Her para hareketinde kategori, ödeme yöntemi, tarih ve fiş/makbuz no kaydı zorunluluğu."),

        ("MODÜL 6: Ürün Alım-Satımı ve Basit Stok Takibi",
         "• Malzeme Kartları: Ürün kodu, ürün adı, birimi, alış fiyatı, satış fiyatı ve asgari stok seviyesi.\n"
         "• Stok Hareketleri: Toptancıdan alış (kasa düşer, stok artar), müşteriye satış (kasa artar, stok düşer), şantiyede sarf malzeme kullanımı.\n"
         "• Kritik Stok Uyarısı: Azalan malzemeler anasayfada otomatik kırmızı uyarı verir; malzeme eksikliği işi durdurmaz.\n"
         "• Ticaret Kârlılığı: Alım-satım yapılan malzemelerden elde edilen kâr ayrıca raporlanır."),

        ("MODÜL 7: Yönetim Dashboard ve Finansal Raporlama",
         "• Şirket Net Kârı: Net Kâr = (Proje Kârları + Malzeme Satış Kârı) - Şirket Sabit Giderleri formülüyle anlık hesaplanır.\n"
         "• Bekleyen Alacak & Borç Raporu: Vadesi gelmiş/gecikmiş müşteri alacakları ve ustalara olan borçlar listesi.\n"
         "• Excel / CSV Dışa Aktarım: Tüm tablolar tek tıkla Excel formatında indirilebilir.")
    ]

    for title, desc in modules:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_before = Pt(8)
        p_m.paragraph_format.space_after = Pt(2)
        r_mt = p_m.add_run(title)
        r_mt.font.bold = True
        r_mt.font.size = Pt(11)
        r_mt.font.color.rgb = ACCENT_COLOR
        
        p_md = doc.add_paragraph()
        p_md.paragraph_format.space_before = Pt(0)
        p_md.paragraph_format.space_after = Pt(6)
        p_md.paragraph_format.line_spacing = 1.15
        r_mdesc = p_md.add_run(desc)
        r_mdesc.font.size = Pt(9.5)
        r_mdesc.font.color.rgb = TEXT_DARK

    # 4. TEKNİK MİMARİ VE KARŞILAŞTIRMA TABLOSU
    h3 = doc.add_heading("3. TEKNİK MİMARİ VE GELİŞTİRİCİYE SIFIR BAĞIMLILIK", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR

    comp_table = doc.add_table(rows=6, cols=2)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_table.autofit = False
    comp_table.columns[0].width = Inches(3.2)
    comp_table.columns[1].width = Inches(3.7)

    headers = [("Piyasadaki Standart / Ağır Yazılımlar", "Sizin İçin Tasarlanan RENOVATION OS")]
    rows_data = [
        ("Ayrı MySQL/Postgre sunucusu gerektirir, şifre/port değişirse çöker.", 
         "Zero-Config SQLite: Veritabanı tek bir güvenli dosyadadır. Sıfır ayar, sıfır port sorunu."),
        ("Her ay kullanıcı başı yüksek lisans ve abonelik ücreti ödetir.", 
         "Sıfır Lisans: Kodlar tamamen firmanıza aittir; ömür boyu sınırsız ve ücretsiz kullanım."),
        ("Yedek almak için teknik yazılımcıya bağımlı kalınır.", 
         "1-Tık Yedek: Veritabanı dosyasını USB belleğe kopyalamanız tam yedek için yeterlidir."),
        ("Karmaşık kurulum ve sunucu ayarları gerektirir.", 
         "1-Tık Kurulum: Klasörü cPanel veya herhangi bir hostinge atıp anında çalıştırma garantisi."),
        ("Ağır arayüzler, mobilde zor kullanılan hantal yapılar.", 
         "Ultra-Hızlı PHP 8.3 + Tailwind: 0.1 saniyede açılan, mobilde sahada kolayca kullanılan modern arayüz.")
    ]

    # Başlık Satırı
    c_th0 = comp_table.cell(0, 0)
    c_th1 = comp_table.cell(0, 1)
    set_cell_background(c_th0, "F1F5F9")
    set_cell_background(c_th1, "0F172A")
    set_cell_margins(c_th0, 100, 100, 120, 120)
    set_cell_margins(c_th1, 100, 100, 120, 120)
    
    r_th0 = c_th0.paragraphs[0].add_run("Piyasadaki Standart / Ağır Sistemler")
    r_th0.font.bold = True
    r_th0.font.size = Pt(9.5)
    r_th0.font.color.rgb = PRIMARY_COLOR
    
    r_th1 = c_th1.paragraphs[0].add_run("RENOVATION OS (Size Özel Çözüm)")
    r_th1.font.bold = True
    r_th1.font.size = Pt(9.5)
    r_th1.font.color.rgb = RGBColor(255, 255, 255)

    # Veri Satırları
    for idx, (col0, col1) in enumerate(rows_data, start=1):
        c_r0 = comp_table.cell(idx, 0)
        c_r1 = comp_table.cell(idx, 1)
        bg = "FFFFFF" if idx % 2 == 1 else "F8FAFC"
        set_cell_background(c_r0, bg)
        set_cell_background(c_r1, "EFF6FF" if idx % 2 == 1 else "E0F2FE")
        set_cell_margins(c_r0, 80, 80, 120, 120)
        set_cell_margins(c_r1, 80, 80, 120, 120)
        
        r0 = c_r0.paragraphs[0].add_run(col0)
        r0.font.size = Pt(9)
        r0.font.color.rgb = TEXT_DARK
        
        r1 = c_r1.paragraphs[0].add_run(col1)
        r1.font.size = Pt(9)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)

    # 5. TİCARİ ŞARTLAR, TESLİMAT VE FİYAT
    h4 = doc.add_heading("4. TİCARİ ŞARTLAR VE TESLİMAT PROTOKOLÜ", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR

    terms_p = doc.add_paragraph()
    terms_p.paragraph_format.line_spacing = 1.2
    terms_p.paragraph_format.space_after = Pt(8)
    
    terms = [
        ("Sistem Geliştirme Bedeli:", " 1.000 USD (veya TCMB efektif satış kuru karşılığı TL)"),
        ("Ödeme Planı:", " %50 Peşinat (500 USD) başlangıç onayı ile; %50 Teslimatta (500 USD) sistem test edilip onaylandıktan sonra."),
        ("Teslim Süresi:", " Sipariş ve peşinat onayından itibaren tam 7 GÜN (1 Hafta)."),
        ("Ara Bilgilendirme:", " 3. günde sistemin canlı önizlemesi (Proje, Müşteri ve Keşif ekranları) onayınıza sunulacaktır."),
        ("Teslim Edilecekler:", " Açık kaynak PHP/SQLite kodları (ZIP), hazır demo veritabanı, resimli kullanım kılavuzu ve yedekleme araçları."),
        ("Garanti & Destek:", " Teslimat sonrası 30 gün boyunca her türlü kullanıcı sorusu ve teknik destek ücretsizdir.")
    ]
    
    for label, val in terms:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(2)
        r_l = p_t.add_run(f"• {label}")
        r_l.font.bold = True
        r_l.font.size = Pt(9.5)
        r_l.font.color.rgb = PRIMARY_COLOR
        r_v = p_t.add_run(val)
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = TEXT_DARK

    # 6. SÖZLEŞME VE ONAY İMZA BÖLÜMÜ
    h5 = doc.add_heading("5. ONAY VE İMZA", level=1)
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(8)
    for r in h5.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_COLOR

    p_appr = doc.add_paragraph()
    p_appr.paragraph_format.space_after = Pt(14)
    r_appr = p_appr.add_run(
        "Bu şartname ve teklif belgesi; yukarıda maddelenen tüm fonksiyonel modülleri, mimariyi ve koşulları eksiksiz karşılayan anahtar teslim bir yazılım taahhüdüdür."
    )
    r_appr.font.size = Pt(9.5)
    r_appr.font.italic = True
    r_appr.font.color.rgb = TEXT_MUTED

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.columns[0].width = Inches(3.45)
    sig_table.columns[1].width = Inches(3.45)
    
    s0 = sig_table.cell(0, 0)
    s1 = sig_table.cell(0, 1)
    set_cell_background(s0, "F8FAFC")
    set_cell_background(s1, "F8FAFC")
    set_cell_margins(s0, 120, 120, 140, 140)
    set_cell_margins(s1, 120, 120, 140, 140)
    
    p_s0 = s0.paragraphs[0]
    p_s0.paragraph_format.space_after = Pt(2)
    r_s0_h = p_s0.add_run("HİZMET SAĞLAYICI (YAZILIM / MİMARLIK)\n\n")
    r_s0_h.font.bold = True
    r_s0_h.font.size = Pt(9.5)
    r_s0_h.font.color.rgb = PRIMARY_COLOR
    p_s0.add_run("Yetkili: Proje Yöneticisi\nİmza: ____________________\nTarih: 25.09.2026\n").font.size = Pt(9)
    
    p_s1 = s1.paragraphs[0]
    p_s1.paragraph_format.space_after = Pt(2)
    r_s1_h = p_s1.add_run("HİZMET ALAN (MÜŞTERİ ONAYI)\n\n")
    r_s1_h.font.bold = True
    r_s1_h.font.size = Pt(9.5)
    r_s1_h.font.color.rgb = PRIMARY_COLOR
    p_s1.add_run("Yetkili İsim / Soyisim:\nİmza: ____________________\nŞirket Kaşesi:\n").font.size = Pt(9)

    doc.save(output_path)
    print(f"Kusursuz Word Belgesi Olusturuldu: {output_path}")

if __name__ == "__main__":
    out_file = "/Users/macair1/projects/html/renovation-os/RENOVATION_OS_MUSTERI_TEKLIF_SARTNAME.docx"
    create_proposal_docx(out_file)
